# -*- coding: utf-8 -*-
"""A conferência de uma operação em PDF — a tela de resultado num documento.

A arquitetura que estes testes protegem: a conta acontece uma vez, no
processamento, e vira um retrato da operação. A tela, a aba
CONFERÊNCIA_AUTOMAÇÃO dentro da planilha e este PDF apenas desenham esse
retrato. Se as três pudessem divergir num centavo, quem assina o processo não
teria como saber qual delas está certa — por isso há um teste que confere
número por número entre a planilha e o PDF.

Depois disso, o que mais importa:

  * a reconciliação impressa FECHA: somar as linhas da folha tem de dar o
    total lido, senão quem confere à mão acusa uma divergência inexistente;
  * o veredito (bate / não bate) se lê antes de qualquer número;
  * uma sessão gravada por uma versão anterior do app ainda gera documento;
  * nada de PII e nada que a marcação do reportlab possa reformatar.
"""
import io
import re
from decimal import Decimal

import pytest
from openpyxl import Workbook

from services import conferencia, conferencia_pdf
from services.utils import formatar_moeda


def _rec(**ajustes):
    base = {
        "total_lido": Decimal("124530.55"),
        "total_preenchido": Decimal("98120.30"),
        "total_estrutura": Decimal("1000.00"),
        "total_fora_escopo": Decimal("20410.25"),
        "total_folha_fora_escopo": Decimal("2000.00"),
        "total_sem_vinculo": Decimal("1500.00"),
        "total_folha_sem_vinculo": Decimal("1000.00"),
        "total_setor_nao_mapeado": Decimal("500.00"),
        "diferenca": Decimal("0.00"),
        "confere": True,
    }
    base.update(ajustes)
    return base


# O retrato de uma operação, como `processar` o grava na sessão.
_RETRATO = {
    "competencia": "07/2026",
    "perfil_nome": "Secretaria de Cultura",
    "arquivo_origem": "ListagemEventos.xlsx",
    "arquivo_modelo": "molde_padrao.xlsx",
    "output_nome": "RETENCAO_PREENCHIDA_20260805_141815.xlsx",
    "aba_destino": "SEC CULTURA",
    "qtd_celulas": 7,
    "reconciliacao": _rec(),
    "por_setor": {"ACS": Decimal("50000.30"), "ADMINISTRATIVO": Decimal("48120.00")},
    "por_coluna": {"CEF": Decimal("58120.30"), "IR": Decimal("40000.00")},
    "por_tipo": {"13º salário": Decimal("10000.00"), "Mensal": Decimal("88120.30")},
    "por_evento": {"IRRF": Decimal("40000.00")},
    "decisoes_folhas": [
        {"rotulo": "MENSAL", "tipo": "Mensal", "total": Decimal("88120.30"),
         "status": "ok", "motivo": "nome igual ao da linha"},
        {"rotulo": "RESCISÃO", "tipo": None, "total": Decimal("1000.00"),
         "status": "sem_vinculo", "motivo": "nenhuma linha corresponde"},
    ],
    "decisoes_rubricas": [
        {"rotulo": "12/120 - EMPRESTIMO CEF 1", "coluna": "CEF",
         "total": Decimal("58120.30"), "status": "regra", "motivo": "regra CEF"},
    ],
    "lotacoes_nao_mapeadas": {"5.9999.0000 - LOTAÇÃO NOVA": 3},
    "rubricas_sem_vinculo": {},
    "rubricas_fora_escopo": {"INSS": 12},
    "folhas_sem_vinculo": {"RESCISÃO": 1},
    "folhas_sugeridas": {"COMPLEMENTAR": 2},
    "folhas_fora_escopo": {},
    "pendencias_estrutura": [
        {"motivo": "linha de tipo não encontrada", "setor": "ACS", "tipo": "Férias",
         "rubrica": "IR", "valor": Decimal("1000.00")},
    ],
}


def _texto(retrato=None) -> str:
    import pypdfium2 as pdfium

    bytes_pdf = conferencia_pdf.montar_pdf(retrato or _RETRATO)
    doc = pdfium.PdfDocument(io.BytesIO(bytes_pdf))
    return "\n".join(pagina.get_textpage().get_text_range() for pagina in doc)


# ===========================================================================
# O documento
# ===========================================================================

def test_e_um_pdf_valido_e_paginado():
    import pypdfium2 as pdfium

    bytes_pdf = conferencia_pdf.montar_pdf(_RETRATO)
    assert bytes_pdf.startswith(b"%PDF-")
    assert len(pdfium.PdfDocument(io.BytesIO(bytes_pdf))) >= 2


def test_traz_as_secoes_da_tela_de_resultado():
    texto = _texto()
    for titulo in ("CONFERÊNCIA DA AUTOMAÇÃO DE RETENÇÕES", "RECONCILIAÇÃO",
                   "TOTAL POR SETOR", "TOTAL POR RUBRICA (COLUNA DA PLANILHA)",
                   "TOTAL POR LINHA DE TIPO DE FOLHA", "COMO CADA FOLHA FOI DESTINADA",
                   "COMO CADA EVENTO FOI DESTINADO", "PENDÊNCIAS E OBSERVAÇÕES"):
        assert titulo in texto, f"seção ausente: {titulo}"


def test_identifica_a_operacao_que_esta_conferindo():
    """Um documento de conferência sem procedência não se confere."""
    texto = _texto()
    for esperado in ("07/2026", "Secretaria de Cultura", "SEC CULTURA",
                     "ListagemEventos.xlsx", "molde_padrao.xlsx",
                     "RETENCAO_PREENCHIDA_20260805_141815.xlsx"):
        assert esperado in texto


def test_numera_as_paginas_e_declara_a_origem():
    texto = _texto()
    assert "Página 1" in texto
    assert "CONFERÊNCIA_AUTOMAÇÃO" in texto  # aponta para o registro exaustivo


# ===========================================================================
# A reconciliação — o número que sustenta o documento
# ===========================================================================

def test_as_linhas_da_reconciliacao_fecham_no_total_lido():
    """Somar a folha impressa tem de dar o total lido.

    `total_preenchido` já vem descontado do que não achou lugar na planilha;
    omitir essa linha faria a soma dar menos e quem confere à mão acusaria
    uma divergência que não existe.
    """
    rec = _RETRATO["reconciliacao"]
    parcelas = sum(rec[chave] for _r, chave, _d in conferencia.LINHAS_RECONCILIACAO
                   if chave != "total_lido")
    assert parcelas == rec["total_lido"]

    texto = _texto()
    for _rotulo, chave, _destaque in conferencia.LINHAS_RECONCILIACAO:
        assert formatar_moeda(rec[chave]) in texto, f"{chave} não aparece no PDF"


def test_veredito_de_divergencia_aparece_com_o_valor():
    texto = _texto(dict(_RETRATO, reconciliacao=_rec(
        diferenca=Decimal("12.34"), confere=False)))
    assert "DIVERGÊNCIA" in texto
    assert formatar_moeda(Decimal("12.34")) in texto


def test_veredito_de_conferencia_ok():
    assert "CONFERE" in _texto()


def test_pdf_e_aba_de_conferencia_nao_podem_divergir_num_centavo():
    """O teste que sustenta a arquitetura das três superfícies.

    A aba dentro da planilha e este PDF leem o MESMO retrato. Se um número
    aparecer só num deles, a separação entre calcular e desenhar foi rompida.
    """
    wb = Workbook()
    conferencia.criar_aba_conferencia(wb, _RETRATO)
    ws = wb["CONFERÊNCIA_AUTOMAÇÃO"]

    numeros = {celula.value for linha in ws.iter_rows() for celula in linha
               if isinstance(celula.value, float)}
    texto = _texto()
    for numero in numeros:
        assert formatar_moeda(Decimal(str(numero))) in texto, \
            f"{numero} está na planilha e não no PDF"


def test_totais_por_dimensao_aparecem_item_a_item():
    texto = _texto()
    for chave in ("por_setor", "por_coluna", "por_tipo"):
        soma = sum(_RETRATO[chave].values(), Decimal("0.00"))
        for rotulo, valor in _RETRATO[chave].items():
            assert rotulo in texto
            assert formatar_moeda(valor) in texto
        assert formatar_moeda(soma) in texto, f"TOTAL de {chave} não aparece"


# ===========================================================================
# As decisões e as pendências
# ===========================================================================

def test_mostra_como_cada_folha_e_cada_evento_foram_destinados():
    texto = _texto()
    assert "MENSAL" in texto and "nome igual ao da linha" in texto
    assert "12/120 - EMPRESTIMO CEF 1" in texto and "CEF" in texto
    # O status vem por extenso, o mesmo rótulo da aba de conferência.
    assert conferencia.ROTULO_STATUS["regra"] in texto


def test_folha_sem_destino_e_dita_por_extenso():
    assert "não preenchida" in _texto()


def test_lista_as_pendencias_com_a_contagem():
    texto = _texto()
    assert "5.9999.0000 - LOTAÇÃO NOVA" in texto
    assert "Itens sem lugar na planilha" in texto
    assert "linha de tipo não encontrada" in texto


def test_sem_pendencias_diz_isso_em_vez_de_omitir_a_secao():
    """Seção ausente parece esquecimento; 'nenhuma' é informação."""
    limpo = dict(_RETRATO, lotacoes_nao_mapeadas={}, rubricas_sem_vinculo={},
                 folhas_sem_vinculo={}, folhas_sugeridas={}, rubricas_fora_escopo={},
                 folhas_fora_escopo={}, pendencias_estrutura=[])
    texto = _texto(limpo)
    assert "PENDÊNCIAS E OBSERVAÇÕES" in texto
    assert "Sem pendências acionáveis" in texto


# ===========================================================================
# Resiliência — a sessão atravessa versões do app
# ===========================================================================

def test_retrato_de_versao_anterior_ainda_gera_documento():
    """Uma sessão gravada antes de um campo existir não pode derrubar o PDF."""
    antigo = {"aba_destino": "ABA", "reconciliacao": {"total_lido": "1000.00",
                                                      "confere": True}}
    assert conferencia_pdf.montar_pdf(antigo).startswith(b"%PDF-")


@pytest.mark.parametrize("estrutura", [None, {}, "lixo", 42])
def test_pendencia_de_estrutura_com_forma_errada_e_ignorada(estrutura):
    """Retrato de outra versão pode trazer a lista noutro formato: a seção
    some, o documento continua — nunca uma exceção no meio da geração."""
    texto = _texto(dict(_RETRATO, pendencias_estrutura=estrutura))
    assert "PENDÊNCIAS E OBSERVAÇÕES" in texto
    assert "Itens sem lugar na planilha" not in texto


def test_retrato_vazio_nao_derruba_a_geracao():
    assert conferencia_pdf.montar_pdf({}).startswith(b"%PDF-")


@pytest.mark.parametrize("lixo", ["corrompido", None, "", [], {"a": 1}])
def test_valor_ilegivel_vira_zero_em_vez_de_excecao(lixo):
    retrato = dict(_RETRATO, reconciliacao=_rec(total_lido=lixo),
                   por_setor={"ACS": lixo})
    assert conferencia_pdf.montar_pdf(retrato).startswith(b"%PDF-")


def test_lista_gigante_e_cortada_com_aviso_em_vez_de_paginar_sem_fim():
    """O PDF é para arquivar; o registro exaustivo é a aba da planilha."""
    grandes = [
        {"rotulo": f"EVENTO {i:05d}", "coluna": "CEF", "total": Decimal("1.00"),
         "status": "ok", "motivo": "teste"}
        for i in range(conferencia_pdf.MAX_LINHAS + 40)
    ]
    texto = _texto(dict(_RETRATO, decisoes_rubricas=grandes))
    assert "EVENTO 00000" in texto
    assert f"Mais 40 linha(s) não cabem" in texto
    assert "CONFERÊNCIA_AUTOMAÇÃO" in texto  # diz onde está a lista inteira


# ===========================================================================
# Guardas de segurança
# ===========================================================================

_VENENOS = [
    "<font size=99>quebra",
    "</para><b>reformata",
    "=cmd|'/c calc'!A1",
    "<script>alert(1)</script>",
    "&entity; & <",
]


@pytest.mark.parametrize("veneno", _VENENOS)
def test_marcacao_em_texto_nao_derruba_nem_reformata_o_documento(veneno):
    """`Paragraph` recebe mini-XML: um "<font size=99>" solto num nome de
    setor derrubaria a geração ou reformataria o documento que alguém assina."""
    retrato = dict(
        _RETRATO, competencia=veneno, perfil_nome=veneno, aba_destino=veneno,
        arquivo_origem=veneno, arquivo_modelo=veneno, output_nome=veneno,
        por_setor={veneno: Decimal("10.00")}, por_coluna={veneno: Decimal("10.00")},
        por_tipo={veneno: Decimal("10.00")},
        decisoes_folhas=[{"rotulo": veneno, "tipo": veneno, "total": Decimal("10.00"),
                          "status": "ok", "motivo": veneno}],
        decisoes_rubricas=[{"rotulo": veneno, "coluna": veneno, "total": Decimal("10.00"),
                            "status": "ok", "motivo": veneno}],
        lotacoes_nao_mapeadas={veneno: 1},
        pendencias_estrutura=[{"motivo": veneno, "setor": veneno, "tipo": veneno,
                               "rubrica": veneno, "valor": Decimal("1.00")}],
    )
    texto = _texto(retrato)                       # não pode levantar
    assert "CONFERÊNCIA DA AUTOMAÇÃO" in texto    # o documento seguiu inteiro


def test_nao_carrega_dados_pessoais():
    """A conferência mostra agregados; nome, matrícula e CPF ficam fora."""
    conteudo = conferencia_pdf.montar_pdf(_RETRATO).decode("latin-1")
    assert not re.search(r"\d{3}\.\d{3}\.\d{3}-\d{2}", conteudo)
    for campo in ("funcionario", "matricula"):
        assert campo not in conteudo.lower()


# ===========================================================================
# Nome do arquivo
# ===========================================================================

def test_nome_do_arquivo_herda_o_carimbo_da_planilha():
    """Os dois arquivos da mesma operação ficam lado a lado na pasta."""
    assert conferencia_pdf.nome_arquivo("RETENCAO_PREENCHIDA_20260805_141815.xlsx") == \
        "CONFERENCIA_20260805_141815.pdf"


@pytest.mark.parametrize("entrada", [
    None, "", "sem carimbo.xlsx", "../../etc/passwd",
    'RETENCAO";rm -rf /.xlsx', "planilha\r\nSet-Cookie: a=b.xlsx",
])
def test_nome_do_arquivo_nunca_carrega_o_texto_recebido(entrada):
    """O nome sai de um carimbo extraído por regex, nunca do texto recebido:
    um `output_nome` adulterado não vira caminho nem cabeçalho HTTP."""
    nome = conferencia_pdf.nome_arquivo(entrada)
    assert re.fullmatch(r"CONFERENCIA_\d{8}_\d{6}\.pdf", nome), nome


# ===========================================================================
# Trabalho limitado — a grade vem de um .xlsx ENVIADO
# ===========================================================================

def test_totais_por_dimensao_tambem_param_no_teto():
    """`molde.MAX_SETORES` só limita o molde DESENHADO na interface.

    Um modelo enviado pode trazer milhares de blocos de setor, e sem teto
    aqui a conferência viraria centenas de páginas — trabalho sem limite a
    partir de um arquivo de fora.
    """
    import pypdfium2 as pdfium

    setores = {f"SETOR {i:05d}": Decimal("10.00")
               for i in range(conferencia_pdf.MAX_LINHAS + 200)}
    bytes_pdf = conferencia_pdf.montar_pdf(dict(_RETRATO, por_setor=setores))
    assert len(pdfium.PdfDocument(io.BytesIO(bytes_pdf))) < 40


def test_total_da_dimensao_cortada_continua_somando_tudo():
    """Cortar a lista não pode mudar a conta: o TOTAL da folha é o da tela."""
    setores = {f"SETOR {i:05d}": Decimal("10.00")
               for i in range(conferencia_pdf.MAX_LINHAS + 200)}
    texto = _texto(dict(_RETRATO, por_setor=setores))
    assert formatar_moeda(Decimal("10.00") * len(setores)) in texto   # TOTAL cheio
    assert formatar_moeda(Decimal("10.00") * 200) in texto            # o que ficou de fora
    assert "já estão no TOTAL acima" in texto


def test_geracao_de_um_retrato_pesado_termina_rapido():
    """Orçamento de tempo: a tela de resultado não pode ficar pendurada."""
    import time

    pesado = dict(
        _RETRATO,
        por_setor={f"SETOR {i:04d}": Decimal("10.00") for i in range(3000)},
        por_coluna={f"RUBRICA {i:04d}": Decimal("10.00") for i in range(3000)},
        decisoes_rubricas=[{"rotulo": f"EVENTO {i:05d}", "coluna": "CEF",
                            "total": Decimal("1.00"), "status": "ok", "motivo": "x"}
                           for i in range(3000)],
        lotacoes_nao_mapeadas={f"LOTAÇÃO {i:04d}": 1 for i in range(3000)},
    )
    inicio = time.monotonic()
    assert conferencia_pdf.montar_pdf(pesado).startswith(b"%PDF-")
    assert time.monotonic() - inicio < 10
